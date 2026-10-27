# -*- coding: utf-8 -*-
"""
把招標文件(須知、需求說明書、評審須知、契約、標價清單)轉成純文字,供後續拆標使用。

支援:.docx .doc .odt .ods .pdf .xlsx .txt .md
用法:
    python -I extract_rfp.py <檔案或資料夾> --out <輸出資料夾>
輸出:每個來源檔一個 .txt,另產生 _index.json(來源路徑、字數、擷取方式、警告)。

注意:
- .doc 需要本機有 Word(用 COM 轉換),沒有就標記警告,請使用者另存 .docx。
- 掃描型 PDF 沒有文字層,會標記 needs_ocr,不要假裝已完整解析。
"""
import argparse, json, os, re, sys, zipfile
from xml.etree import ElementTree as ET

SUPPORTED = {".docx", ".doc", ".odt", ".ods", ".pdf", ".xlsx", ".txt", ".md"}


def read_docx(path):
    from docx import Document
    d = Document(path)
    out = []
    # 依文件順序輸出段落與表格
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            t = "".join(x.text or "" for x in child.iter() if x.tag.endswith("}t"))
            if t.strip():
                out.append(t.strip())
        elif tag == "tbl":
            from docx.table import Table
            tb = Table(child, d)
            for row in tb.rows:
                cells, prev = [], None
                for c in row.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text.replace("\n", " / ").strip())
                out.append(" | ".join(cells))
            out.append("")
    return "\n".join(out), "python-docx", []


def read_doc(path):
    try:
        import win32com.client
    except Exception:
        return "", "none", ["no_word_com: 請另存成 .docx 再處理"]
    import tempfile
    tmp = os.path.join(tempfile.mkdtemp(), "x.docx")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True, AddToRecentFiles=False)
        doc.SaveAs2(tmp, FileFormat=16)
        doc.Close(False)
    finally:
        word.Quit()
    text, _, w = read_docx(tmp)
    return text, "word-com->docx", w


def read_odf(path):
    ns = {"t": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
          "tb": "urn:oasis:names:tc:opendocument:xmlns:table:1.0"}
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("content.xml"))
    out = []

    def txt(e):
        return "".join(e.itertext()).strip()
    for e in root.iter():
        tag = e.tag.split("}")[1]
        if tag in ("p", "h") and not any(True for _ in []):
            pass
    # 簡化:逐列輸出表格,其餘段落照順序
    body = root.find(".//{urn:oasis:names:tc:opendocument:xmlns:office:1.0}body")
    def walk(node):
        for ch in node:
            tag = ch.tag.split("}")[1]
            if tag in ("p", "h"):
                t = txt(ch)
                if t:
                    out.append(t)
            elif tag == "table":
                for row in ch.iter("{%s}table-row" % ns["tb"]):
                    cells = [txt(c) for c in row if c.tag.endswith("table-cell")]
                    if any(cells):
                        out.append(" | ".join(cells))
                out.append("")
            else:
                walk(ch)
    walk(body)
    return "\n".join(out), "odf-xml", []


def read_pdf(path):
    warns = []
    text = ""
    try:
        import pdfplumber
        pages = []
        with pdfplumber.open(path) as pdf:
            for i, p in enumerate(pdf.pages, 1):
                pages.append(f"--- 第 {i} 頁 ---\n" + (p.extract_text() or ""))
        text = "\n".join(pages)
    except Exception as e:
        warns.append(f"pdfplumber_failed: {e}")
    if len(re.sub(r"--- 第 \d+ 頁 ---|\s", "", text)) < 50:
        warns.append("needs_ocr: PDF 幾乎沒有文字層(可能是掃描檔),不可宣稱已完整解析")
    return text, "pdfplumber", warns


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== 工作表:{ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out), "openpyxl", []


def extract(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return read_docx(path)
    if ext == ".doc":
        return read_doc(path)
    if ext in (".odt", ".ods"):
        return read_odf(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".xlsx":
        return read_xlsx(path)
    return open(path, encoding="utf-8", errors="replace").read(), "plain", []


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    files = []
    if os.path.isdir(a.src):
        for r, _, fs in os.walk(a.src):
            for f in fs:
                if f.startswith("~$"):
                    continue
                if os.path.splitext(f)[1].lower() in SUPPORTED:
                    files.append(os.path.join(r, f))
    else:
        files = [a.src]
    index = []
    for i, f in enumerate(sorted(files), 1):
        try:
            text, how, warns = extract(f)
        except Exception as e:
            text, how, warns = "", "error", [f"error: {e}"]
        name = f"{i:02d}_" + re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(f)) + ".txt"
        with open(os.path.join(a.out, name), "w", encoding="utf-8") as fh:
            fh.write(text)
        index.append({"source": f, "txt": name, "chars": len(text), "method": how, "warnings": warns})
        print(f"{name}  {len(text)} 字  {how}  {'; '.join(warns)}")
    with open(os.path.join(a.out, "_index.json"), "w", encoding="utf-8") as fh:
        json.dump(index, fh, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    sys.exit(main())

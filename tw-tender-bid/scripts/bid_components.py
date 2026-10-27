# -*- coding: utf-8 -*-
"""
bid_components.py -- 版面元件(建立在 bid_kit 之上)

stats     重點數字卡(例:服務年資、完成案件數、團隊人數)
callout   提示框(note 提醒、key 重點、warn 注意;用左側色條,不用圖示或 Emoji)
steps     步驟流程(大號序號 + 標題 + 說明)
divider   章節扉頁(整頁:大號章號、章名、副標;本身就是 Heading 1,會進目錄)
twocol    雙欄對照(例:現況與做法、需求與回應)
lead      導言段落(章節開頭的一兩句重點)
cover_art 依風格色票自動生成的幾何封面圖,種子取自案名,每案構圖不同

所有元件都讀 bid_kit 目前的風格與規格(顏色、字體、字級),所以換風格或換規格會一起變。
"""
import io, os, zlib

import bid_kit as K
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _cell_text(cell, text, size, color, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, font=None):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text and len(cell.paragraphs) == 1 else cell.add_paragraph()
    p.alignment = align
    K.spc(p, before=before, after=after, lr=K.LINE_RULE, ln=(K.LINE_TW if K.LINE_RULE == "exact" else 280))
    K.run(p, text, bold=bold, size=size, color=color, font=font)
    return p


def _plain_table(doc, rows, cols, widths=None):
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    K.set_table_borders(t, pad=(60, 140))
    if widths:
        t.autofit = False
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def lead(doc, text):
    """導言:比內文大一級、主色,放在章標題之後。"""
    p = doc.add_paragraph()
    K.spc(p, before=40, after=100, lr=K.LINE_RULE, ln=(K.LINE_TW if K.LINE_RULE == "exact" else 360))
    K.run(p, text, size=K.SZ["body"] + 2, color=K.NAVY, font=K.ZH_HEAD)
    return p


def stats(doc, items):
    """items: [{"value": "15", "unit": "年", "label": "影像製作經驗"}, ...](最多 4 欄)。"""
    n = max(1, min(4, len(items)))
    tw = K.text_width_cm()
    t = _plain_table(doc, 1, n, [tw / n] * n)
    for i, it in enumerate(items[:n]):
        c = t.rows[0].cells[i]
        K.set_cell_bg(c, K.PALE)
        K.cell_border(c, top=(24, str(K.ACCENT)), left=(48, "FFFFFF"), right=(48, "FFFFFF"))
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = c.paragraphs[0]
        K.spc(p, before=80, after=0)
        K.run(p, str(it["value"]), bold=True, size=K.SZ["h1"] * 2, color=K.NAVY, font=K.ZH_HEAD)
        if it.get("unit"):
            K.run(p, f" {it['unit']}", size=K.SZ["body"], color=K.GRAY)
        _cell_text(c, it.get("label", ""), K.SZ["body"], K.INK, before=20, after=80)
    doc.add_paragraph()
    return t


_KINDS = {"note": ("TEAL", "tint"), "key": ("ACCENT", "tint"), "warn": ("NAVY", "zebra")}


def callout(doc, title, text, kind="note"):
    bar, bg = _KINDS.get(kind, _KINDS["note"])
    bar_c = getattr(K, bar)
    bg_c = K.PALE if bg == "tint" else K.MIST
    t = _plain_table(doc, 1, 1, [K.text_width_cm()])
    c = t.rows[0].cells[0]
    K.set_cell_bg(c, bg_c)
    K.cell_border(c, left=(48, str(bar_c)))
    if title:
        _cell_text(c, title, K.SZ["body"], bar_c if kind != "note" else K.NAVY, bold=True, before=60, after=20, font=K.ZH_HEAD)
    _cell_text(c, text, K.SZ["body"], K.INK, before=0 if title else 60, after=60)
    doc.add_paragraph()
    return t


def steps(doc, items):
    """items: [{"title": "...", "text": "..."}, ...]。左側大號序號。"""
    tw = K.text_width_cm()
    t = _plain_table(doc, len(items), 2, [2.0, tw - 2.0])
    for i, it in enumerate(items):
        c0, c1 = t.rows[i].cells
        c0.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        _cell_text(c0, f"{i + 1:02d}", K.SZ["h1"] * 1.6, K.ACCENT, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=20, font=K.ZH_HEAD)
        _cell_text(c1, it["title"], K.SZ["h3"], K.NAVY, bold=True, before=40, after=10, font=K.ZH_HEAD)
        if it.get("text"):
            _cell_text(c1, it["text"], K.SZ["body"], K.INK, after=40)
        K.cell_border(c0, bottom=(4, str(K.LINE)))
        K.cell_border(c1, bottom=(4, str(K.LINE)))
    K.table_rows_flow(t, header_rows=0)
    doc.add_paragraph()
    return t


def twocol(doc, left, right):
    """left / right: {"title": "...", "items": ["...", ...]}。"""
    tw = K.text_width_cm()
    t = _plain_table(doc, 2, 2, [tw / 2, tw / 2])
    for ci, side in enumerate((left, right)):
        h, b = t.rows[0].cells[ci], t.rows[1].cells[ci]
        K.set_cell_bg(h, K.NAVY)
        K.cell_border(h, left=(48, "FFFFFF"), right=(48, "FFFFFF"))
        _cell_text(h, side["title"], K.SZ["body"], K.WHITE, bold=True, before=40, after=40, font=K.ZH_HEAD)
        K.set_cell_bg(b, K.PALE)
        K.cell_border(b, left=(48, "FFFFFF"), right=(48, "FFFFFF"))
        first = True
        for it in side["items"]:
            if first:
                p = b.paragraphs[0]
                first = False
            else:
                p = b.add_paragraph()
            K.spc(p, before=30, after=30, lr=K.LINE_RULE, ln=(K.LINE_TW if K.LINE_RULE == "exact" else 290))
            K.run(p, "▪  ", size=K.SZ["body"], color=K.ACCENT)
            K.run(p, it, size=K.SZ["body"], color=K.INK)
    K.table_rows_flow(t, header_rows=1)
    doc.add_paragraph()
    return t


def divider(doc, number, title, subtitle=None, bm=None):
    """章節扉頁。整頁一個章名,本身是 Heading 1(進導覽窗格與目錄)。number 例如 "01" 或 "壹"。"""
    K.strip_trailing_blanks(doc)
    p0 = doc.add_paragraph()
    p0.paragraph_format.page_break_before = True
    K.spc(p0, before=2600, after=0)
    K.run(p0, str(number), bold=True, size=K.SZ["h1"] * 4, color=K.PALE if K.TABLE_STYLE != "minimal" else K.LINE, font=K.ZH_HEAD)
    p = doc.add_paragraph(style="Heading 1")
    K.spc(p, before=0, after=120)
    K.para_border(p, "bottom", 24, K.ACCENT, space=8)
    K.run(p, title, bold=True, size=K.SZ["h1"] * 1.9, color=K.NAVY, font=K.ZH_HEAD)
    K.keep_with_next(p)
    if bm:
        K.add_bookmark(p, bm)
    if subtitle:
        p2 = doc.add_paragraph()
        K.spc(p2, before=120, after=0)
        K.run(p2, subtitle, size=K.SZ["body"] + 1, color=K.GRAY)
    pb = doc.add_paragraph()
    pb.paragraph_format.page_break_before = True
    pb.paragraph_format.space_after = Pt(0)
    return p


# ── 封面幾何圖 ───────────────────────────────────────────────────
def cover_art(kind, seed_text, width_cm, height_cm):
    """依目前風格色票生成透明底 PNG(位元組)。同一案名會得到同一張,不同案名構圖不同。"""
    import random
    import matplotlib.pyplot as plt
    from matplotlib.patches import Circle, Rectangle, Wedge
    rng = random.Random(zlib.crc32((seed_text or "case").encode("utf-8")))
    pal = [K.C_NAVY, K.C_STEEL, K.C_TEAL, K.C_PALE]
    acc = K.C_ACC
    W, H = 100.0, 100.0 * height_cm / width_cm
    fig = plt.figure(figsize=(width_cm / 2.54, height_cm / 2.54))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W); ax.set_ylim(0, H); ax.axis("off")

    if kind == "blocks":
        cols, rows = 14, 4
        cw, ch = W / cols, H / rows
        placed_accent = False
        for r in range(rows):
            for c in range(cols):
                roll = rng.random()
                x, y = c * cw, r * ch
                if roll < 0.30:
                    ax.add_patch(Rectangle((x + .8, y + .8), cw - 1.6, ch - 1.6, color=rng.choice(pal[:3]), lw=0))
                elif roll < 0.42:
                    ax.add_patch(Rectangle((x + .8, y + .8), cw - 1.6, ch - 1.6, fill=False, ec=pal[0], lw=1.0))
                elif roll < 0.52:
                    ax.add_patch(Rectangle((x + .8, y + .8), cw - 1.6, ch - 1.6, color=pal[3], lw=0))
                elif roll < 0.55 and not placed_accent:
                    ax.add_patch(Rectangle((x + .8, y + .8), cw - 1.6, ch - 1.6, color=acc, lw=0)); placed_accent = True
        if not placed_accent:
            ax.add_patch(Rectangle((W - 2 * cw + .8, .8), cw - 1.6, ch - 1.6, color=acc, lw=0))
    elif kind == "rings":
        cx, cy = W * rng.uniform(.62, .82), H * rng.uniform(.35, .65)
        for i, r in enumerate([46, 38, 31, 24, 18, 12, 7]):
            ax.add_patch(Circle((cx, cy), r, fill=False, ec=pal[i % 3], lw=1.0 + (i % 3) * .7))
        ax.add_patch(Circle((cx, cy), 3.2, color=acc, lw=0))
        ax.add_patch(Wedge((cx, cy), 46, 180, 270, width=3, color=pal[3], lw=0))
    elif kind == "bars":
        n = 30
        base = rng.uniform(.18, .3)
        for i in range(n):
            h = H * min(.97, base + .62 * (i / n) ** 1.2 + rng.uniform(-.08, .08))
            level = min(3, int(h / H * 4))
            col = [pal[3], pal[2], pal[1], pal[0]][level]
            ax.add_patch(Rectangle((i * W / n + .5, 0), W / n - 1.0, h, color=col, lw=0))
        top = max(range(n), key=lambda k: k)
        ax.add_patch(Rectangle((top * W / n + .5, 0), W / n - 1.0, H * .98, color=acc, lw=0))
    elif kind == "waves":
        import numpy as np
        x = np.linspace(0, W, 400)
        for i in range(9):
            amp = H * (.10 + .035 * i)
            freq = rng.uniform(.035, .06)
            ph = rng.uniform(0, 6.28)
            y = H * .5 + amp * np.sin(freq * x + ph) * (1 - .35 * np.cos(x / W * 3.14))
            ax.plot(x, y, color=(acc if i == 6 else pal[i % 3]), lw=(1.8 if i == 6 else 0.9))
    elif kind == "dots":
        import numpy as np
        cols, rows = 46, max(6, int(46 * H / W))
        slope = rng.uniform(.25, .6)
        for r in range(rows):
            for c in range(cols):
                d = abs((c / cols) * slope - (r / rows)) / max(slope, .5)
                rad = max(.0, 1.25 - 2.6 * d)
                if rad > .12:
                    ax.add_patch(Circle(((c + .5) * W / cols, (r + .5) * H / rows), rad * W / cols * .55,
                                        color=(acc if (r + c) % 29 == 0 else pal[0]), lw=0))
    else:
        plt.close(fig)
        return None
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, transparent=True)
    plt.close(fig)
    return buf.getvalue()

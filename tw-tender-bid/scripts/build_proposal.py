# -*- coding: utf-8 -*-
"""
把 spec(JSON)產生成服務建議書 / 企劃書 Word 檔。

用法:
    python -I build_proposal.py <spec.json> --out <輸出.docx> [--seals] [--finalize]

spec 格式見 references/spec-format.md,範例見 assets/spec_example.json。
--seals     插入簽名與印章(未經使用者逐份確認不要加)
--finalize  用本機 Word 更新所有欄位(目錄頁碼、頁碼、章名、圖表編號)並存檔,
            這樣交付的 docx 開啟時不會跳出「是否更新欄位」的提示。需要 Windows 與 Word。
"""
import argparse, json, os, re, sys, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bid_kit as K  # noqa: E402
import bid_components as C  # noqa: E402
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.shared import Cm  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402

_CN_DIGITS = "零一二三四五六七八九"
_CN_BIG = "零壹貳參肆伍陸柒捌玖"
_STRIP = re.compile(r"^(?:[壹貳參肆伍陸柒捌玖拾]+、|[一二三四五六七八九十]+、|[（(][一二三四五六七八九十]+[）)]|\d+[.、)）]\s*|\(\d+\)\s*)")
_CAP_STRIP = re.compile(r"^\s*[圖表]\s*\d+\s*[:：]?\s*")
_LEVEL_OF = {"banner": 0, "divider": 0, "subh": 1, "subsubh": 2}


def _cn(n, digits):
    """1 到 99 的中文數字。"""
    if n < 10:
        return digits[n]
    tens, ones = divmod(n, 10)
    unit = "拾" if digits is _CN_BIG else "十"
    t = unit if tens == 1 else digits[tens] + unit
    return t + (digits[ones] if ones else "")


def _label(scheme, n):
    if scheme == "cn_big":
        return _cn(n, _CN_BIG) + "、"
    if scheme == "cn_small":
        return _cn(n, _CN_DIGITS) + "、"
    if scheme == "cn_paren":
        return "（" + _cn(n, _CN_DIGITS) + "）"
    if scheme == "arabic":
        return f"{n}."
    if scheme == "arabic_paren":
        return f"({n})"
    raise ValueError(f"未知的標號格式:{scheme}")


def apply_numbering(sections, schemes):
    """依招標文件規定的標號格式,替章節標題加標號(先移除原有標號)。
    schemes 例:["cn_big", "cn_small", "cn_paren", "arabic", "arabic_paren"],對應第 1 到 5 層;
    banner 與 divider 為第 1 層、subh 為第 2 層、subsubh 為第 3 層。"""
    counters = [0] * 5
    for s in sections:
        lv = _LEVEL_OF.get(s["type"])
        if lv is None or lv >= len(schemes):
            continue
        counters[lv] += 1
        for j in range(lv + 1, 5):
            counters[j] = 0
        title_key = "title" if s["type"] == "divider" else "text"
        s[title_key] = _label(schemes[lv], counters[lv]) + _STRIP.sub("", s[title_key])
        s["_numbered"] = True
        if s["type"] == "subh":
            s["num"] = ""
    return sections


def _chart(doc, c):
    kind = c.get("kind")
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    cols = K.CHART_COLORS
    if kind == "pie":
        labels = [d["label"] for d in c["data"]]
        vals = [d["value"] for d in c["data"]]
        ax.pie(vals, labels=labels, autopct="%1.1f%%", startangle=90, colors=cols[:len(vals)],
               wedgeprops={"linewidth": 1, "edgecolor": "white"})
        ax.axis("equal")
    elif kind == "bar":
        labels = [d["label"] for d in c["data"]]
        vals = [d["value"] for d in c["data"]]
        ax.bar(labels, vals, color=cols[1])
        for i, v in enumerate(vals):
            ax.text(i, v, f"{v:,}", ha="center", va="bottom", fontsize=8)
        ax.spines[["top", "right"]].set_visible(False)
    elif kind == "gantt":
        tasks = c["data"]
        for i, t in enumerate(tasks):
            ax.barh(i, t["end"] - t["start"], left=t["start"], color=cols[i % 3], height=0.55)
        ax.set_yticks(range(len(tasks)))
        ax.set_yticklabels([t["task"] for t in tasks], fontsize=9)
        ax.invert_yaxis()
        ax.set_xlabel(c.get("unit", "週"))
        ax.grid(axis="x", alpha=0.3)
        ax.spines[["top", "right"]].set_visible(False)
    else:
        plt.close(fig)
        raise ValueError(f"未知的 chart kind: {kind}")
    cap = _CAP_STRIP.sub("", c["caption"]) if c.get("caption") else None
    K.chart_img(doc, fig, w=c.get("width", 15.0), cap=cap, label="圖" if cap else None)


def _toc(doc, items):
    """目錄:章名、點線引導、頁碼(PAGEREF 欄位)。items: [(層級, 文字, 書籤)]。"""
    K.banner(doc, "目錄", sub="Table of Contents", new_page=True)
    tw = Cm(K.text_width_cm())
    for level, name, bm in items:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(tw, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        K.spc(p, before=60 if level == 0 else 10, after=20)
        if level:
            p.paragraph_format.left_indent = Cm(0.8)
        K.run(p, name, bold=(level == 0), size=K.SZ["body"],
              color=K.NAVY if level == 0 else K.INK, font=K.ZH_HEAD if level == 0 else None)
        p.add_run("\t")
        K.add_pageref(p, bm, size=K.SZ["body"], color=K.NAVY, bold=(level == 0))


def build(spec, out, seals=False):
    case = spec.get("case", {})
    fmt = spec.get("format")
    if isinstance(fmt, str):  # format.json 的路徑
        with open(fmt, encoding="utf-8") as f:
            fmt = json.load(f)
    K.configure(
        case_title=case.get("title", ""), case_short=case.get("short", ""),
        client=case.get("client", ""), subtitle=case.get("subtitle", ""),
        cover_date=case.get("cover_date", ""), sign_date=case.get("sign_date", ""),
        partner=case.get("partner"), blind=bool(case.get("blind", False)),
        style=spec.get("style") or case.get("theme"), format=fmt,
        cover_image=case.get("cover_image"),
    )
    doc = K.new_doc()
    sections = spec["sections"]
    if (fmt or {}).get("numbering"):
        apply_numbering(sections, fmt["numbering"])

    # 目錄:章(banner、divider)與節(subh)都給書籤
    toc_levels = int(spec.get("toc_levels", 2))
    toc_items, nb, ns = [], 0, 0
    for s in sections:
        if s["type"] in ("banner", "divider"):
            nb += 1
            s.setdefault("bm", f"sec{nb}")
            toc_items.append((0, s["text"] if s["type"] == "banner" else s["title"], s["bm"]))
        elif s["type"] == "subh" and toc_levels >= 2:
            ns += 1
            s.setdefault("bm", f"sub{ns}")
            toc_items.append((1, s["text"], s["bm"]))

    if spec.get("cover", True):
        K.build_cover(doc, doc_kind=spec.get("doc_kind", "服務建議書"), doc_sub=case.get("doc_sub"))
    if spec.get("toc", True) and toc_items:
        _toc(doc, toc_items)

    first_banner = True
    for s in sections:
        t = s["type"]
        if t == "banner":
            K.banner(doc, s["text"], sub=s.get("sub"), bm=s["bm"],
                     new_page=s.get("new_page", True) if not first_banner or spec.get("cover", True) else False)
            first_banner = False
        elif t == "divider":
            C.divider(doc, s.get("number", ""), s["title"], s.get("subtitle"), bm=s["bm"])
            first_banner = False
        elif t == "subh":
            K.subh(doc, s["text"], num=(s.get("num") if not s.get("_numbered") else ""),
                   new_page=s.get("new_page", False), bm=s.get("bm"))
        elif t == "subsubh":
            K.subsubh(doc, s["text"], marker="" if s.get("_numbered") else "▸  ")
        elif t == "body":
            K.body(doc, s["text"], indent=s.get("indent"))
        elif t == "lead":
            C.lead(doc, s["text"])
        elif t == "bullets":
            for it in s["items"]:
                K.bullet(doc, it)
        elif t == "numlist":
            K.numlist_block(doc, s["items"], start=s.get("start", 1))
        elif t == "table":
            if s.get("caption"):
                K.caption(doc, _CAP_STRIP.sub("", s["caption"]), label="表")
            K.tbl(doc, s["headers"], s["rows"], cw=s.get("widths"), size=s.get("size"),
                  center_from=s.get("center_from"))
        elif t == "kv":
            K.kv_table(doc, [tuple(p) for p in s["pairs"]])
        elif t == "hbox":
            K.hbox(doc, s["text"])
        elif t == "quote":
            K.quotebox(doc, s.get("title", ""), s["text"])
        elif t == "callout":
            C.callout(doc, s.get("title", ""), s["text"], kind=s.get("kind", "note"))
        elif t == "stats":
            C.stats(doc, s["items"])
        elif t == "steps":
            C.steps(doc, s["items"])
        elif t == "twocol":
            C.twocol(doc, s["left"], s["right"])
        elif t == "image":
            if os.path.exists(s["path"]):
                p = doc.add_paragraph()
                p.alignment = 1
                p.add_run().add_picture(s["path"], width=Cm(s.get("width", 14)))
                if s.get("caption"):
                    K.caption(doc, _CAP_STRIP.sub("", s["caption"]), label="圖")
            else:
                K.body(doc, f"【待補:圖檔不存在 {s['path']}】")
        elif t == "chart":
            _chart(doc, s)
        elif t == "page_break":
            K.page_break(doc)
        elif t == "sign_block":
            K.sign_block(doc, label_date=s.get("date", True), seals=seals)
        else:
            raise ValueError(f"未知的 section type: {t}")

    K.add_header_footer(doc, header_text=case.get("short") or None, skip_first=bool(spec.get("cover", True)))
    K.set_update_fields_on_open(doc)
    K.strip_trailing_blanks(doc)
    K.save_doc(doc, out)
    return out


def finalize(path):
    """用 Word 更新所有欄位並存檔,再移除「開啟時更新欄位」的設定,避免開檔時跳提示。"""
    try:
        import win32com.client
    except Exception:
        print("未安裝 pywin32 或不是 Windows,略過 finalize(開檔時 Word 會詢問是否更新欄位)")
        return False
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(path))
        d.Fields.Update()
        for sec in d.Sections:
            for hf in (sec.Headers(1), sec.Footers(1)):
                hf.Range.Fields.Update()
        d.Save()
        d.Close(False)
    finally:
        word.Quit()
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                data = re.sub(rb"<w:updateFields[^>]*/>", b"", data)
            zout.writestr(item, data)
    os.replace(tmp, path)
    return True


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--out", required=True)
    ap.add_argument("--seals", action="store_true", help="插入簽名與印章(須經使用者確認)")
    ap.add_argument("--finalize", action="store_true", help="用 Word 更新欄位並存檔")
    a = ap.parse_args()
    with open(a.spec, encoding="utf-8") as f:
        spec = json.load(f)
    print("產生:", build(spec, a.out, seals=a.seals))
    if a.finalize:
        print("finalize:", "完成" if finalize(a.out) else "略過")


if __name__ == "__main__":
    main()
